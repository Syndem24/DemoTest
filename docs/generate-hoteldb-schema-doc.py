"""Generate docs/HotelDb-Schema.docx + HotelDb-Schema.md from the live Mori Hotel schema.

Single source of truth for the database documentation. Run:  py -3 generate-hoteldb-schema-doc.py
Data verified against sys.tables/sys.columns/sys.foreign_keys on HotelDb (LocalDB).
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x0D, 0x94, 0x88)
OUT_DIR = Path(__file__).resolve().parent

# ----------------------------------------------------------------------------
# Schema data — (table, entity kind, purpose, [(column, type, key/null, notes)])
# ----------------------------------------------------------------------------
TABLES = [
    ("AccountUser", "Strong entity",
     "One login account per user — hotel staff (AdminManager, Receptionist) and guests who sign in with Google. ASP.NET Identity user store.",
     [
        ("Id", "nvarchar(450)", "PK", "Identity GUID string"),
        ("UserName / NormalizedUserName", "nvarchar(256)", "", "Sign-in name + normalized lookup"),
        ("Email / NormalizedEmail", "nvarchar(256)", "", "Account email + normalized lookup"),
        ("EmailConfirmed", "bit", "", "Email ownership confirmed"),
        ("PasswordHash", "nvarchar(max)", "", "Local password hash — guests get one via forced first-login setup"),
        ("SecurityStamp / ConcurrencyStamp", "nvarchar(max)", "", "Identity security + concurrency stamps"),
        ("PhoneNumber / PhoneNumberConfirmed", "nvarchar(max) / bit", "", "Optional phone + confirmation"),
        ("TwoFactorEnabled", "bit", "", "Identity 2FA flag (unused — Google flow instead)"),
        ("LockoutEnd / LockoutEnabled / AccessFailedCount", "datetimeoffset / bit / int", "", "Identity lockout controls"),
        ("MustChangePassword", "bit", "", "Forces /Account/ChangePassword on next login"),
        ("FullName", "nvarchar(120)", "", "Display name (staff employee name or guest name)"),
        ("Address", "nvarchar(300)", "", "Optional address"),
        ("BirthDate", "date", "NULL", "Optional birth date"),
        ("GoogleEmail / NormalizedGoogleEmail", "nvarchar(256)", "NULL, unique idx", "Linked Gmail for sign-in/recovery"),
        ("GoogleVerificationStatus", "nvarchar(40)", "", "NotLinked | PendingGoogleVerification | GoogleVerified"),
        ("RoleId", "nvarchar(450)", "FK → AccountRole, NULL", "Single role per account"),
        ("DashboardLayoutJson", "nvarchar(max)", "NULL, JSON", "Staff dashboard GridStack layout"),
     ]),
    ("AccountRole", "Strong entity (lookup)",
     "Application roles: AdminManager, Receptionist, Guest.",
     [
        ("Id", "nvarchar(450)", "PK", "Identity GUID string"),
        ("Name / NormalizedName", "nvarchar(256)", "", "Role name + normalized lookup"),
        ("ConcurrencyStamp", "nvarchar(max)", "", "Identity concurrency stamp"),
     ]),
    ("AccountExternalLogin", "Weak entity (composite PK)",
     "External OAuth logins linked to an account — currently Google guest sign-in.",
     [
        ("LoginProvider", "nvarchar(450)", "PK (composite)", "'Google'"),
        ("ProviderKey", "nvarchar(450)", "PK (composite)", "Provider's user id"),
        ("ProviderDisplayName", "nvarchar(max)", "NULL", "Display name from provider"),
        ("UserId", "nvarchar(450)", "FK → AccountUser, CASCADE", "Owning account"),
     ]),
    ("AccountAuthToken", "Weak entity (composite PK)",
     "ASP.NET Identity tokens per account (refresh/recovery artifacts).",
     [
        ("UserId", "nvarchar(450)", "PK (composite), FK → AccountUser CASCADE", "Owning account"),
        ("LoginProvider", "nvarchar(450)", "PK (composite)", "Token issuer"),
        ("Name", "nvarchar(450)", "PK (composite)", "Token name"),
        ("Value", "nvarchar(max)", "NULL", "Token payload"),
     ]),
    ("PasswordResetCode", "Weak entity",
     "One-time hashed 6-digit reset codes emailed to any account — staff and Google guests alike "
     "(guests get a local PasswordHash via forced first-login setup). ForgotPassword has no role "
     "restriction.",
     [
        ("Id", "int", "PK, identity", ""),
        ("UserId", "nvarchar(450)", "", "Target account id (logical link, no FK)"),
        ("NormalizedEmail", "nvarchar(256)", "", "Email the code was sent to"),
        ("CodeHash", "nvarchar(128)", "", "Hashed code — never stored plain"),
        ("CreatedAtUtc / ExpiresAtUtc", "datetime2", "", "Issue + expiry"),
        ("ConsumedAtUtc", "datetime2", "NULL", "Set when used"),
        ("FailedAttempts", "int", "", "Brute-force counter"),
     ]),
    ("Booking", "Strong entity (aggregate root)",
     "One guest stay — online booking, walk-in, or OTA channel. Holds guest contact inline (guests need no account to book).",
     [
        ("Id", "int", "PK, identity", ""),
        ("Reference", "nvarchar(24)", "unique", "Public confirmation code"),
        ("GuestName / GuestEmail / GuestPhone", "nvarchar(120/254/40)", "", "Guest contact (inline, no Customer table)"),
        ("Kind", "nvarchar(20)", "", "Booking | Reservation (lead-time split)"),
        ("Status", "nvarchar(20)", "", "Pending | Confirmed | Rejected | Cancelled | CheckedOut"),
        ("Channel", "nvarchar(30)", "", "Online | WalkIn | FrontDeskExtension | Agoda | Expedia | RedDoorz | OtherThirdParty"),
        ("PaymentOption", "nvarchar(20)", "", "Full | Half due at booking"),
        ("CheckInAtUtc / CheckoutTimeUtc", "datetime2", "NULL col, required by app", "Scheduled stay window"),
        ("TotalAmount", "decimal(18,2)", "", "Room nights + charges (locked at booking rate)"),
        ("AmountDueNow", "decimal(18,2)", "", "Amount due at booking time"),
        ("AdultCount / ChildCount", "int", "", "Head counts"),
        ("GuestPartyJson", "nvarchar(4000)", "NULL, JSON", "Per-room head-count breakdown"),
        ("SpecialOfferId", "int", "FK → SpecialOffer, SET NULL", "Promo applied at booking"),
        ("CashOnlyPromo", "bit", "", "Walk-in limited-time promo → cash only"),
        ("ArrivalDiscountRequest", "nvarchar(30)", "", "None | SeniorCitizen | Pwd"),
        ("ArrivalWarningSentAtUtc / PendingCallWarningSentAtUtc / CheckoutWarningSentAtUtc", "datetime2", "NULL", "20-min warning timestamps"),
        ("IsNotificationCleared", "bit", "", "Hidden from admin bell"),
        ("IsArchived / ArchivedAtUtc", "bit / datetime2", "", "Soft-delete to history"),
        ("CreatedAtUtc / UpdatedAtUtc", "datetime2", "", "Audit timestamps"),
     ]),
    ("BookingItem", "Weak entity",
     "One room-type line on a stay (qty x nightly rate). Rates are locked at booking time; promo/regular never repriced afterwards.",
     [
        ("Id", "int", "PK, identity", ""),
        ("BookingId", "int", "FK → Booking, CASCADE", "Owning stay"),
        ("RoomTypeId", "int", "FK → RoomType, SET NULL", "Nullable — keeps RoomTypeName if type deleted"),
        ("RoomTypeName", "nvarchar(100)", "", "Name snapshot at booking time"),
        ("Quantity", "int", "", "Rooms of this type"),
        ("PricePerNight", "decimal(18,2)", "", "Rate used (promo or regular)"),
     ]),
    ("BookingRoomAssignment", "Associative entity (M:N junction)",
     "Joins one physical Room to a BookingItem — reception assigns door numbers.",
     [
        ("Id", "int", "PK, identity", ""),
        ("BookingItemId", "int", "FK → BookingItem, CASCADE", ""),
        ("RoomId", "int", "FK → Room, NO ACTION", ""),
     ]),
    ("BookingCharge", "Weak entity",
     "Extra stay fee line — early check-in, late checkout, extra person, incidental, service fee, snack/beverage, stay extension, arrival discount, loyalty coupon.",
     [
        ("Id", "int", "PK, identity", ""),
        ("BookingId", "int", "FK → Booking, CASCADE", "Owning stay"),
        ("ChargeType", "nvarchar(30)", "", "EarlyCheckIn | LateCheckout | ExtraPerson | Incidental | ServiceFee | SnackBeverage | StayExtension | ArrivalDiscount | LoyaltyCoupon"),
        ("Label", "nvarchar(200)", "", "Display label"),
        ("Quantity / Nights", "int", "", "Rooms/hours/persons x night multiplier"),
        ("UnitAmount / Amount", "decimal(18,2)", "", "Rate and total (StayExtension excluded from fee sum; LoyaltyCoupon may be negative)"),
        ("CreatedAtUtc", "datetime2", "", ""),
     ]),
    ("PaymentRecord", "Strong entity (append-only ledger)",
     "One payment/void event per stay. Corrections are voids, never deletes. Receipt-photo/OCR pipeline removed — payments post manually.",
     [
        ("Id", "int", "PK, identity", ""),
        ("BookingId", "int", "FK → Booking, CASCADE", "Owning stay"),
        ("ReceiptNumber", "nvarchar(40)", "", "Issued receipt no."),
        ("EventType", "nvarchar(30)", "", "Deposit | ArrivalPayment | BalanceSettlement | Refund | Adjustment"),
        ("Method", "nvarchar(30)", "", "Cash | EWallet | Other (+ legacy Card/BankTransfer/Maya)"),
        ("Amount", "decimal(18,2)", "", "Positive payment; negative = refund/adjustment"),
        ("StayTotalAtPosting / BalanceAfter", "decimal(18,2)", "", "Ledger snapshots at posting"),
        ("PaidAtUtc", "datetime2", "", "When collected"),
        ("ReceivedBy", "nvarchar(120)", "", "Staff collector name (snapshot string, not FK)"),
        ("ExternalReference / BankTransferReference", "nvarchar(120)", "NULL", "E-wallet/InstaPay ref (manual entry); bank ref legacy"),
        ("Status", "nvarchar(20)", "", "Posted | Voided"),
        ("VerifiedAtUtc / VerifiedBy", "datetime2 / nvarchar(120)", "NULL", "Staff manual verification of e-wallet payment"),
        ("VoidedAtUtc / VoidReason / VoidedBy", "datetime2 / nvarchar(500) / nvarchar(120)", "NULL", "Void details"),
        ("Notes", "nvarchar(1000)", "NULL", "Free text (e.g., cash change given)"),
     ]),
    ("Room", "Strong entity",
     "One physical guest room (door number).",
     [
        ("Id", "int", "PK, identity", ""),
        ("RoomTypeId", "int", "FK → RoomType, NO ACTION", "Category"),
        ("RoomNumber", "nvarchar(20)", "unique", "Door number"),
        ("Status", "nvarchar(20)", "", "Available | Unavailable | Occupied | Cleaning (shown as Maintaining)"),
     ]),
    ("RoomType", "Strong entity",
     "Sellable room category (Queen, Twin, ...) with shared rate, occupancy, photos, inclusions.",
     [
        ("RoomTypeId", "int", "PK, identity", ""),
        ("Name", "nvarchar(100)", "unique", "Category name"),
        ("Description", "nvarchar(max)", "NULL", "Marketing copy"),
        ("PricePerNight", "decimal(18,2)", "", "Regular nightly rate"),
        ("MaxOccupancy / BedCount", "int", "", "Capacity"),
        ("Inclusions / Images", "nvarchar(max)", "JSON", "Standard inclusions list + photo paths"),
        ("CreatedAt", "datetime2", "", ""),
     ]),
    ("SpecialOffer", "Strong entity",
     "Promo rate for one RoomType (Limited Time, Stay Longer, Google Loyalty). Sibling rows share a campaign title across room types.",
     [
        ("Id", "int", "PK, identity", ""),
        ("RoomTypeId", "int", "FK → RoomType, CASCADE", "Discounted category"),
        ("Kind", "nvarchar(40)", "", "LimitedTime | StayLongerSaveMore | GoogleLoyalty (+ legacy kinds)"),
        ("Title / Description", "nvarchar(160) / nvarchar(1000)", "", "Campaign name + copy"),
        ("RegularPricePerNight", "decimal(18,2)", "", "'Was' price shown crossed out"),
        ("PromoPricePerNight", "decimal(18,2)", "NULL", "Promo rate (null for informational kinds)"),
        ("MinNights", "int", "NULL", "StayLongerSaveMore threshold (>= 2)"),
        ("Channels", "int", "flags", "OnlineVisible | WalkIn | FrontDesk | ThirdPartyVisible"),
        ("CashOnly", "bit", "", "Stay payments must be cash"),
        ("LoyaltyApplyMode", "int", "", "EveryNight | FirstNight | WeeklyReset | FirstBooking"),
        ("StartsAtUtc / EndsAtUtc / OpenEnded", "datetime2 / datetime2 / bit", "", "Validity window"),
        ("IsActive / SortOrder", "bit / int", "", "Live flag + display order"),
        ("CreatedAtUtc / UpdatedAtUtc", "datetime2", "", ""),
     ]),
    ("StayReview", "Weak entity (1:0..1 of Booking)",
     "One verified post-checkout review per booking — unique index on BookingId enforces the cap.",
     [
        ("Id", "int", "PK, identity", ""),
        ("BookingId", "int", "FK → Booking, CASCADE, unique", "Reviewed stay"),
        ("GuestUserId", "nvarchar(450)", "indexed, no FK", "AccountUser.Id of reviewer (logical link)"),
        ("DisplayName", "nvarchar(80)", "", "Shown publicly"),
        ("OverallRating / StaffRating / ComfortRating / FacilitiesRating", "tinyint", "", "1-5 stars"),
        ("WouldRecommend", "bit", "NULL", ""),
        ("Comment / TagsJson", "nvarchar(2000) / nvarchar(1000)", "NULL / JSON", "Text + tag chips"),
        ("IsPublished", "bit", "", "Visible on guest site"),
        ("HotelReply / HotelReplyAtUtc / HotelReplyBy", "nvarchar(1000) / datetime2 / nvarchar(120)", "NULL", "Staff reply"),
        ("HasHotelReply", "bit", "computed (persisted)", "Derived flag for fast reply filtering"),
        ("CreatedAtUtc / UpdatedAtUtc", "datetime2", "", ""),
     ]),
    ("SecureSetting", "Strong entity (config vault)",
     "Encrypted configuration — SMTP password, Google OAuth credentials, Gemini/Groq keys.",
     [
        ("Id", "int", "PK, identity", ""),
        ("Key", "nvarchar(80)", "unique", "Setting key (SecureSettingKeys)"),
        ("Ciphertext", "nvarchar(max)", "", "Encrypted value"),
        ("UpdatedUtc", "datetime2", "", ""),
     ]),
    ("SystemAuditLog", "Strong entity (log)",
     "Integrity trail: who / what / when / target / why for account, payment, booking, offer, config, file, shift, review actions.",
     [
        ("Id", "bigint", "PK, identity", ""),
        ("AtUtc", "datetime2", "", "When"),
        ("Intent / Domain / Action", "nvarchar(40/40/80)", "", "AdministrativeAction|ConfigurationChange|FileModification; Payment|Account|Booking|..."),
        ("ActorUserId / ActorDisplayName", "nvarchar(450) / nvarchar(120)", "", "Who did it"),
        ("TargetType / TargetId / TargetLabel", "nvarchar(40/80/200)", "", "What was touched (e.g. TargetType='AccountUser')"),
        ("Reason / Summary", "nvarchar(500)", "NULL / required", "Why + human summary"),
     ]),
    ("SystemFlushLog", "Strong entity (log)",
     "Audit trail for export-then-delete operations (booking history, payments, staff audit) — what was flushed, when, by whom, into which file.",
     [
        ("Id", "int", "PK, identity", ""),
        ("Kind", "nvarchar(40)", "", "BookingHistory | Payments | StaffAudit"),
        ("FlushedAtUtc / PerformedBy", "datetime2 / nvarchar(120)", "", "When + who"),
        ("RecordCount / FileName / Summary", "int / nvarchar(200) / nvarchar(2000)", "", "What was exported"),
     ]),
]

# (child, fk column, parent, cardinality, delete rule, meaning)
RELATIONSHIPS = [
    ("BookingItem", "BookingId", "Booking", "1 : N", "CASCADE",
     "A stay has 1+ room-type lines; deleting a booking removes its lines."),
    ("BookingCharge", "BookingId", "Booking", "1 : N", "CASCADE",
     "A stay has 0+ fee lines."),
    ("PaymentRecord", "BookingId", "Booking", "1 : N", "CASCADE",
     "A stay has 0+ posted payment events (append-only)."),
    ("StayReview", "BookingId", "Booking", "1 : 0..1", "CASCADE",
     "At most one review per stay — unique index on BookingId."),
    ("Booking", "SpecialOfferId", "SpecialOffer", "0..1 : N", "SET NULL",
     "A stay may carry one promo; deleting the offer keeps the booking (rate stays locked)."),
    ("BookingItem", "RoomTypeId", "RoomType", "0..1 : N", "SET NULL",
     "Line keeps RoomTypeName snapshot if the type is deleted."),
    ("BookingRoomAssignment", "BookingItemId", "BookingItem", "1 : N", "CASCADE",
     "Each line gets 0+ physical room assignments."),
    ("BookingRoomAssignment", "RoomId", "Room", "1 : N", "NO ACTION",
     "A physical room can appear on many lines over time."),
    ("Room", "RoomTypeId", "RoomType", "1 : N", "NO ACTION",
     "A type owns many physical rooms."),
    ("SpecialOffer", "RoomTypeId", "RoomType", "1 : N", "CASCADE",
     "Sibling offer rows per type share a campaign title."),
    ("AccountUser", "RoleId", "AccountRole", "0..1 : N", "SET NULL",
     "Each account has at most one role; deleting a role unassigns accounts."),
    ("AccountExternalLogin", "UserId", "AccountUser", "1 : N", "CASCADE",
     "An account can link several external logins (Google)."),
    ("AccountAuthToken", "UserId", "AccountUser", "1 : N", "CASCADE",
     "Identity tokens per account."),
    ("StayReview", "GuestUserId", "AccountUser", "0..1 : N (logical)", "no FK",
     "Indexed string link — enforced by app, not a constraint."),
    ("PasswordResetCode", "UserId", "AccountUser", "1 : N (logical)", "no FK",
     "Reset codes reference accounts by id string."),
    ("PaymentRecord", "ReceivedBy / VerifiedBy / VoidedBy", "AccountUser", "—", "no FK",
     "Staff names stored as snapshot strings, deliberately not FKs."),
]

# Legend tables
ENTITY_KINDS = [
    ("Strong entity", "Has its own PK and independent existence",
     "AccountUser, AccountRole, Booking, Room, RoomType, SpecialOffer, PaymentRecord, SecureSetting, SystemAuditLog, SystemFlushLog"),
    ("Weak entity", "Existence depends on a parent (composite or parent-owned PK)",
     "BookingItem, BookingCharge, StayReview, AccountExternalLogin, AccountAuthToken, PasswordResetCode"),
    ("Associative (junction) entity", "Resolves a many-to-many between two entities",
     "BookingRoomAssignment (BookingItem x Room)"),
]

ATTRIBUTE_KINDS = [
    ("Primary key (simple)", "One column identifies the row", "Booking.Id, Room.Id, AccountUser.Id"),
    ("Primary key (composite)", "Several columns together form the key", "AccountExternalLogin (LoginProvider+ProviderKey), AccountAuthToken (UserId+LoginProvider+Name)"),
    ("Foreign key", "References a parent PK; may be optional (NULL)", "Booking.SpecialOfferId, BookingItem.RoomTypeId, AccountUser.RoleId"),
    ("Simple attribute", "Single atomic value", "Room.RoomNumber, Booking.TotalAmount"),
    ("JSON-valued attribute", "Structured list/object serialized in one column", "RoomType.Inclusions, RoomType.Images, Booking.GuestPartyJson, StayReview.TagsJson"),
    ("Enum-as-string attribute", "Constrained set stored as text", "Booking.Status, BookingChannel, PaymentRecord.Method, Room.Status"),
    ("Flags enum (bitmask int)", "Combinable options packed into one int", "SpecialOffer.Channels"),
    ("Derived / computed attribute", "Value computed by SQL, stored for indexing", "StayReview.HasHotelReply"),
    ("Audit attributes", "Created/updated/sent timestamps", "CreatedAtUtc, UpdatedAtUtc, *WarningSentAtUtc, PaidAtUtc"),
    ("Snapshot attributes", "Copied at write time so later edits don't rewrite history", "BookingItem.RoomTypeName, PaymentRecord.ReceivedBy, StayTotalAtPosting, BalanceAfter"),
]

RELATIONSHIP_KINDS = [
    ("One-to-many (1 : N)", "Parent row has many children; FK on child", "Booking → BookingItem / BookingCharge / PaymentRecord; RoomType → Room / SpecialOffer"),
    ("One-to-zero-or-one (1 : 0..1)", "At most one child per parent; unique FK index", "Booking → StayReview"),
    ("Many-to-many via junction (M : N)", "Associative table with FKs to both sides", "BookingItem ↔ Room through BookingRoomAssignment"),
    ("Optional many-to-one (N : 0..1)", "Nullable FK — child survives parent deletion", "Booking → SpecialOffer (SET NULL), BookingItem → RoomType (SET NULL)"),
    ("Logical link (no FK)", "Id/name stored as string; integrity enforced by app", "StayReview.GuestUserId, PasswordResetCode.UserId, PaymentRecord.ReceivedBy"),
]

CARDINALITY_NOTES = [
    ("1 : N", "One parent, many children — most relationships", "1 stay has N payment records"),
    ("1 : 0..1", "One parent, at most one child", "1 stay has at most 1 review"),
    ("N : 0..1", "Many children may optionally point to one parent", "N stays may reference 1 special offer (or none)"),
    ("M : N", "Many-to-many, resolved by a junction table", "BookingItem ↔ Room via BookingRoomAssignment"),
    ("CASCADE", "Delete parent → children deleted", "Booking deleted → items, charges, payments, review deleted"),
    ("SET NULL", "Delete parent → child FK becomes NULL, row kept", "SpecialOffer deleted → Booking.SpecialOfferId cleared"),
    ("NO ACTION", "Delete blocked while children exist", "RoomType can't be deleted while Rooms reference it"),
]

# ----------------------------------------------------------------------------
# Emitters
# ----------------------------------------------------------------------------

def _md_escape(s):
    return s.replace("|", "\\|")

def write_markdown():
    lines = [
        "# Mori International Hotel — Database Schema",
        "",
        "Generated from the live EF Core model (`HotelBookingDbContext`) and verified against",
        "`sys.tables` / `sys.columns` / `sys.foreign_keys` on `HotelDb`. SQL Server (LocalDB),",
        "dates stored in UTC, shown in Manila time.",
        "",
        "17 tables. Identity tables are named `Account*` (`AccountUser`, `AccountRole`,",
        "`AccountExternalLogin`, `AccountAuthToken`) — they hold **staff and Google guests**.",
        "`PasswordResetCode` serves both staff and guests (every guest sets a local password on first",
        "login).",
        "",
        "## 1. Table catalog — what each entity is for",
        "",
        "| Table | Entity kind | Purpose |",
        "|---|---|---|",
    ]
    for name, kind, purpose, _ in TABLES:
        lines.append(f"| `{name}` | {_md_escape(kind)} | {_md_escape(purpose)} |")
    lines += ["", "---", "", "## 2. Entities and attributes", ""]
    for name, kind, purpose, cols in TABLES:
        lines += [f"### `{name}` — {kind}", "", purpose, "",
                  "| Attribute | Type | Key / Null | Notes |", "|---|---|---|---|"]
        for col, typ, key, note in cols:
            lines.append(f"| `{col}` | {typ} | {_md_escape(key)} | {_md_escape(note)} |")
        lines.append("")
    lines += ["---", "", "## 3. Relationships and cardinality", "",
              "| Child table | FK column | Parent table | Cardinality | On delete | Meaning |",
              "|---|---|---|---|---|---|"]
    for child, col, parent, card, delete, meaning in RELATIONSHIPS:
        lines.append(f"| `{child}` | `{col}` | `{parent}` | {card} | {delete} | {_md_escape(meaning)} |")
    lines += ["", "---", "", "## 4. Legend — entity / attribute / relationship / cardinality types used",
              "", "### Entity types", "", "| Type | Meaning | In this database |", "|---|---|---|"]
    for t, m, ex in ENTITY_KINDS:
        lines.append(f"| {t} | {_md_escape(m)} | {_md_escape(ex)} |")
    lines += ["", "### Attribute types", "", "| Type | Meaning | Examples |", "|---|---|---|"]
    for t, m, ex in ATTRIBUTE_KINDS:
        lines.append(f"| {t} | {_md_escape(m)} | {_md_escape(ex)} |")
    lines += ["", "### Relationship types", "", "| Type | Meaning | Examples |", "|---|---|---|"]
    for t, m, ex in RELATIONSHIP_KINDS:
        lines.append(f"| {t} | {_md_escape(m)} | {_md_escape(ex)} |")
    lines += ["", "### Cardinality & delete rules", "", "| Notation | Meaning | Example here |", "|---|---|---|"]
    for t, m, ex in CARDINALITY_NOTES:
        lines.append(f"| {t} | {_md_escape(m)} | {_md_escape(ex)} |")
    lines.append("")
    (OUT_DIR / "HotelDb-Schema.md").write_text("\n".join(lines), encoding="utf-8")


def set_font(run, size=10, bold=False, color=NAVY):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_font(r, size=20 if level == 1 else 15 if level == 2 else 12, bold=True,
                 color=NAVY if level < 3 else TEAL)


def body(doc, text, size=10):
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=size)
    p.paragraph_format.space_after = Pt(6)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def grid(doc, headers, rows, size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(h), size=size, bold=True)
        shade(cell, "D9EDED")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(val), size=size, color=RGBColor(0x22, 0x22, 0x22))
    return table


def write_docx():
    doc = Document()
    heading(doc, "Mori International Hotel — Database Schema", 1)
    body(doc, "Entity–attribute–relationship reference generated from the live EF Core model "
              "(HotelBookingDbContext) and verified against sys.tables/sys.columns/sys.foreign_keys "
              "on HotelDb (SQL Server / LocalDB). Dates are stored in UTC and displayed in Manila time.")
    body(doc, "17 tables. The Identity tables are named Account* (AccountUser, AccountRole, "
              "AccountExternalLogin, AccountAuthToken) because they hold staff AND Google guests. "
              "PasswordResetCode likewise serves both — every guest sets a local password on "
              "first login.")

    heading(doc, "1. Table catalog — what each entity is for", 2)
    grid(doc, ["Table", "Entity kind", "Purpose"],
         [(n, k, p) for n, k, p, _ in TABLES], size=8)

    heading(doc, "2. Entities and attributes", 2)
    for name, kind, purpose, cols in TABLES:
        heading(doc, f"{name}  ({kind})", 3)
        body(doc, purpose, size=9)
        grid(doc, ["Attribute", "Type", "Key / Null", "Notes"],
             [(c, t, k, n) for c, t, k, n in cols], size=8)

    heading(doc, "3. Relationships and cardinality", 2)
    grid(doc, ["Child table", "FK column", "Parent table", "Cardinality", "On delete", "Meaning"],
         RELATIONSHIPS, size=8)

    heading(doc, "4. Legend — entity / attribute / relationship / cardinality types used", 2)
    heading(doc, "Entity types", 3)
    grid(doc, ["Type", "Meaning", "In this database"], ENTITY_KINDS, size=8)
    heading(doc, "Attribute types", 3)
    grid(doc, ["Type", "Meaning", "Examples"], ATTRIBUTE_KINDS, size=8)
    heading(doc, "Relationship types", 3)
    grid(doc, ["Type", "Meaning", "Examples"], RELATIONSHIP_KINDS, size=8)
    heading(doc, "Cardinality & delete rules", 3)
    grid(doc, ["Notation", "Meaning", "Example here"], CARDINALITY_NOTES, size=8)

    doc.save(OUT_DIR / "HotelDb-Schema.docx")


if __name__ == "__main__":
    write_markdown()
    write_docx()
    print("Wrote HotelDb-Schema.md and HotelDb-Schema.docx")
